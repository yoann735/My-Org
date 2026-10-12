#!/usr/bin/env python3
"""
MealWeek — import du classeur « Plans repas » vers src/mealweek/data/mealweek_data.json

Usage (depuis la racine du dépôt) :
    pip install openpyxl            # une fois
    python3 scripts/mealweek-import/importer_classeur.py [chemin.xlsx]

Par défaut, lit `Plans_repas_9_semaines_-_24_recettes_-_Chronodrive_Compiegne.xlsx`
à la racine du dépôt.

IDEMPOTENT : le JSON est entièrement régénéré à partir du classeur, dans un
ordre déterministe, avec des clés stables (R01…R24, S1…S9, OW01…). Relancer
sur le même classeur produit un fichier identique octet pour octet ; relancer
après avoir modifié le classeur resynchronise tout (rien n'est dupliqué :
chaque recette n'existe qu'une fois, sous sa clé, et les semaines y réfèrent
par id).

FORMULES : le classeur calcule nutrition, totaux de courses, kcal et macros
du planning par formules. On ne lit JAMAIS les valeurs en cache : tout est
recalculé ici à partir des données saisies (quantités, valeurs /100 g, prix).
Les valeurs en cache (data_only) ne servent qu'au contrôle de cohérence
imprimé en fin d'exécution.

Sources :
- « Recettes v3 (1 pers) »  → ingrédients (base 1 portion, substitutions appliquées)
- « Étapes (2 portions) »   → méta (source, temps, ustensiles, substitutions) + étapes
- « Nutrition »             → nutrition par portion (recalculée ingrédient par ingrédient)
- « Sommaire »              → composition des semaines, coûts annoncés (contrôle)
- « S1 » … « S9 »           → planning + liste de courses + placard + week-end
- « Options week-end »      → catalogue des plats surgelés / pizzas
- « Méthode »               → texte des règles
"""
import hashlib
import json
import os
import re
import sys
from datetime import date

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("openpyxl manquant : pip install openpyxl")

RACINE = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
XLSX_DEFAUT = os.path.join(RACINE, 'Plans_repas_9_semaines_-_24_recettes_-_Chronodrive_Compiegne.xlsx')
SORTIE = os.path.join(RACINE, 'src', 'mealweek', 'data', 'mealweek_data.json')

BUDGET = 55
REPAS_PAR_SEMAINE = 14

avertissements = []


def avertir(msg):
    avertissements.append(msg)


# ------------------------------------------------------------------ utilitaires
def lignes(ws):
    """toutes les lignes de la feuille (valeurs brutes, formules en texte)"""
    return [list(r) for r in ws.iter_rows(values_only=True)]


def txt(v):
    return '' if v is None else str(v).strip()


def num(v):
    """nombre ou None (les formules « =… » et le texte donnent None)"""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return v
    return None


def arrondi(v, n=2):
    if v is None:
        return None
    r = round(v + 0.0, n)
    return int(r) if r == int(r) else r


def fmt_qte(v):
    """130 -> '130', 0.5 -> '0.5' (même rendu que la V2)"""
    if v is None:
        return ''
    r = arrondi(v, 3)
    return str(r)


def rid(n):
    return 'R%02d' % int(n)


def titre_recette(cell):
    """'10. Côtes d'échine…' -> (10, "Côtes d'échine…") ; sinon None"""
    m = re.match(r'^\s*(\d{1,2})\.\s+(.+?)\s*$', txt(cell))
    return (int(m.group(1)), m.group(2)) if m else None


def ids_recettes(texte):
    """'5, 15' / '9' / 5 -> ['R05', 'R15']"""
    return [rid(n) for n in re.findall(r'\d+', txt(texte))]


def non_vides(row):
    return [c for c in row if c is not None and txt(c) != '']


# ------------------------------------------------------------------ Étapes (2 portions)
def lire_etapes(ws):
    """{num: {nom, source, temps, ustensiles, substitutions, ing2: [...], etapes: [...]}}"""
    out = {}
    cur = None
    mode = None
    for row in lignes(ws):
        a = row[0] if row else None
        vals = non_vides(row)
        if not vals:
            mode = None if mode == 'etapes' else mode
            continue
        t = titre_recette(a) if len(vals) == 1 else None
        if t:
            cur = {'num': t[0], 'nom': t[1], 'ing2': [], 'etapes': []}
            out[t[0]] = cur
            mode = 'meta'
            continue
        if cur is None:
            continue
        cle = txt(a)
        if mode == 'meta' and cle in ('Source', 'Temps', 'Ustensiles', 'Substitutions'):
            cur[cle.lower()] = txt(row[1])
            continue
        if cle.startswith('Ingrédients'):
            mode = 'ing'
            continue
        if cle == 'Étape':
            mode = 'etapes'
            continue
        if mode == 'ing':
            cur['ing2'].append({
                'nom': cle, 'qte': num(row[1]), 'unite': txt(row[2]),
                'equivalent': txt(row[3]), 'produit': txt(row[4]), 'origine': txt(row[5]),
            })
        elif mode == 'etapes' and num(a) is not None:
            cur['etapes'].append({'numero': int(a), 'titre': '', 'texte': txt(row[1])})
    return out


# ------------------------------------------------------------------ Recettes v3 (1 pers)
def lire_v3(ws):
    """{num: {nom, ingredients: [...], notes: set}} dans l'ordre du fichier"""
    out = {}
    entete = None
    for row in lignes(ws):
        if entete is None:
            if txt(row[0]) == '# recette':
                entete = row
            continue
        if num(row[0]) is None:
            continue
        n = int(row[0])
        r = out.setdefault(n, {'nom': txt(row[1]), 'ingredients': [], 'notes': []})
        q = num(row[3])
        unite = txt(row[4])
        r['ingredients'].append({
            'nom': txt(row[2]),
            'qty_1portion': (fmt_qte(q) + (' ' + unite if unite else '')).strip(),
            'qty_valeur': arrondi(q, 3) if q is not None else None,
            'unite': unite,
            'equivalent': txt(row[5]) or None,
            'produit_chronodrive': txt(row[6]),
            'origine': txt(row[7]),
        })
        note = txt(row[8]) if len(row) > 8 else ''
        if note and note not in r['notes']:
            r['notes'].append(note)
    return out


# ------------------------------------------------------------------ Nutrition
CHAMPS_NUT = ('kcal', 'proteines_g', 'glucides_g', 'lipides_g')


def lire_nutrition(ws):
    """{num: {'total': {...}, 'ingredients': [(nom, qte)], 'ligne_total': n°excel}}
    Recalcul : valeur = Qté/portion × (valeur /100) / 100, somme par recette."""
    out = {}
    cur = None
    for i, row in enumerate(lignes(ws), start=1):  # i = n° de ligne Excel
        vals = non_vides(row)
        t = titre_recette(row[0]) if len(vals) == 1 else None
        if t:
            cur = {'num': t[0], 'acc': [0.0, 0.0, 0.0, 0.0], 'ingredients': []}
            out[t[0]] = cur
            continue
        if cur is None or not vals:
            continue
        cle = txt(row[0])
        if cle in ('Ingrédient',):
            continue
        if cle.startswith('TOTAL 1 PORTION'):
            cur['ligne_total'] = i
            continue
        if cle.startswith('x2') or cle.startswith('%'):
            continue
        q = num(row[1])
        if q is None:
            continue
        per100 = [num(row[k]) or 0 for k in (3, 4, 5, 6)]
        for k in range(4):
            cur['acc'][k] += q * per100[k] / 100
        cur['ingredients'].append((cle, q))
    for r in out.values():
        r['total'] = {c: round(v, 1) for c, v in zip(CHAMPS_NUT, r['acc'])}
    return out


# ------------------------------------------------------------------ Options week-end
def lire_options(ws):
    out = []
    entete = False
    for row in lignes(ws):
        if not entete:
            entete = txt(row[0]) == 'Type'
            continue
        if not txt(row[1]):
            continue
        poids, prix, repas = num(row[2]), num(row[3]), num(row[5])
        out.append({
            'id': 'OW%02d' % (len(out) + 1),
            'type': txt(row[0]),
            'categorie': 'pizza' if txt(row[0]).lower().startswith('pizza') else 'plat',
            'produit': txt(row[1]),
            'poids_g': poids,
            'prix': prix,
            'prix_kg': arrondi(prix * 1000 / poids, 2) if prix and poids else None,
            'repas_couverts': repas,
            'prix_repas': arrondi(prix / repas, 2) if prix and repas else None,
        })
    return out


# ------------------------------------------------------------------ Sommaire
def lire_sommaire(ws):
    rows = lignes(ws)
    titre = txt(rows[0][0])
    modele = txt(rows[1][0])
    semaines = {}
    for row in rows:
        k = txt(row[0])
        if not re.match(r'^S\d+$', k):
            continue
        recettes = []
        for morceau in txt(row[1]).split(' · '):
            t = titre_recette(morceau)
            if t:
                recettes.append(rid(t[0]))
            else:
                avertir(f'Sommaire {k} : recette sans numéro « {morceau} »')
        semaines[k] = {
            'recettes': recettes, 'plat': txt(row[2]), 'pizza': txt(row[3]),
            'courses': num(row[4]), 'placard': num(row[5]), 'weekend': num(row[6]),
            'total': num(row[7]), 'marge': num(row[8]), 'gachis': num(row[9]), 'stock': num(row[10]),
        }
    return titre, modele, semaines


# ------------------------------------------------------------------ feuilles S1…S9
def lire_semaine(ws, recettes, options_par_nom):
    rows = lignes(ws)
    titre = txt(rows[0][0])
    planning, courses, weekend, placard = [], [], [], []
    section = None
    for row in rows:
        a = txt(row[0])
        if a == 'Jour':
            section = 'planning'; continue
        if a.startswith('Produit exact (Chronodrive)'):
            section = 'courses'; continue
        if a.startswith('TOTAL COURSES RECETTES'):
            section = 'weekend'; continue
        if a.startswith('TOTAL PANIER'):
            section = None; continue
        if a == 'Produit exact':
            section = 'placard'; continue
        if a.startswith('TOTAL placard'):
            section = None; continue
        if a.startswith('ÉTAPES DES') or a in ('BUDGET', 'LISTE DE COURSES CHRONODRIVE — formats entiers', 'PLANNING'):
            section = None; continue
        if not non_vides(row):
            if section == 'planning':
                section = None
            continue

        if section == 'planning':
            soir = txt(row[2])
            t = titre_recette(soir)
            recette_id = rid(t[0]) if t else None
            kcal, macros = None, None
            if recette_id and recette_id in recettes:
                n = recettes[recette_id]['nutrition_1portion']
                kcal = int(round(n['kcal']))
                macros = ' / '.join(str(int(round(n[c]))) for c in ('proteines_g', 'glucides_g', 'lipides_g'))
            elif recette_id:
                avertir(f'{ws.title} : recette {recette_id} du planning introuvable')
            kcal_libelle = txt(row[3]) if not txt(row[3]).startswith('=') else None
            if kcal is None and kcal_libelle:
                m = re.search(r'\d[\d\s]*', kcal_libelle)
                kcal = int(m.group(0).replace(' ', '')) if m else None
            planning.append({
                'jour': a,
                'midi': txt(row[1]) or None,
                'soir': soir or None,
                'recette_id': recette_id,
                'kcal_portion': kcal,
                'kcal_libelle': kcal_libelle,
                'a_cuisiner': txt(row[4]) or None,
                'note': txt(row[5]) or None,
                'macros': macros,
            })
        elif section == 'courses':
            qte, pu = num(row[2]), num(row[3])
            acheté, besoin = num(row[6]), num(row[5])
            courses.append({
                'produit': a,
                'rayon': txt(row[1]),
                'qte_acheter': qte,
                'prix_unitaire': pu,
                'total': arrondi(qte * pu, 2),
                'besoin_semaine': besoin,
                'contenu_achete': acheté,
                'unite': txt(row[7]),
                'reste': num(row[8]) or 0,
                'gachis_eur': arrondi(num(row[9]) or 0, 4),
                'stock_eur': arrondi(num(row[10]) or 0, 4),
                'recettes': ids_recettes(row[11]),
            })
        elif section == 'weekend':
            qte, pu = num(row[2]), num(row[3])
            opt = options_par_nom.get(a)
            if not opt:
                avertir(f'{ws.title} : « {a} » absent du catalogue Options week-end')
            weekend.append({
                'produit': a,
                'rayon': txt(row[1]),
                'qte': qte,
                'prix_unitaire': pu,
                'total': arrondi(qte * pu, 2),
                'usage': txt(row[11]),
                'option_id': opt['id'] if opt else None,
                'categorie': opt['categorie'] if opt else ('pizza' if 'pizza' in a.lower() else 'plat'),
            })
        elif section == 'placard':
            qte, pf = num(row[2]), num(row[3])
            utilise, contenu = num(row[5]), num(row[6])
            reste = num(row[8])
            placard.append({
                'produit': a,
                'rayon': txt(row[1]),
                'qte': qte,
                'prix_format': pf,
                'total': arrondi(qte * pf, 2),
                'qte_utilisee': utilise,
                'contenu_format': contenu,
                'unite': txt(row[7]),
                'reste': reste,
                'valeur_reste': arrondi(pf * reste / contenu, 4) if pf and reste is not None and contenu else None,
                'recettes': ids_recettes(row[11]),
            })
    return titre, planning, courses, weekend, placard


# ------------------------------------------------------------------ assemblage
def construire(chemin):
    wb = openpyxl.load_workbook(chemin)               # formules en texte : on recalcule
    cache = openpyxl.load_workbook(chemin, data_only=True)  # contrôle uniquement

    etapes = lire_etapes(wb['Étapes (2 portions)'])
    v3 = lire_v3(wb['Recettes v3 (1 pers)'])
    nut = lire_nutrition(wb['Nutrition'])
    options = lire_options(wb['Options week-end'])
    options_par_nom = {o['produit']: o for o in options}
    titre, modele, sommaire = lire_sommaire(wb['Sommaire'])
    methode = [txt(r[0]) for r in lignes(wb['Méthode']) if txt(r[0])]

    # ---- recettes : une seule fois chacune, clé = numéro ----
    recettes = {}
    nut_brut = {}
    for n in sorted(set(etapes) | set(v3) | set(nut)):
        e, v, nu = etapes.get(n), v3.get(n), nut.get(n)
        if not (e and v and nu):
            avertir(f'Recette {n} incomplète (étapes={bool(e)}, v3={bool(v)}, nutrition={bool(nu)})')
            continue
        if e['nom'] != v['nom']:
            avertir(f'Recette {n} : nom différent entre Étapes « {e["nom"]} » et v3 « {v["nom"]} »')
        # contrôle : quantités Étapes = 2 × v3, et Nutrition = v3
        q1 = [(i['nom'], i['qty_valeur']) for i in v['ingredients']]
        q2 = [(i['nom'], i['qte']) for i in e['ing2']]
        if [x[0] for x in q1] != [x[0] for x in q2]:
            avertir(f'Recette {n} : liste d\'ingrédients différente entre v3 et Étapes')
        else:
            for (nm, a), (_, b) in zip(q1, q2):
                if a is not None and b is not None and abs(a * 2 - b) > 1.01:
                    avertir(f'Recette {n} : {nm} = {a} (v3, 1 pers.) mais {b} (Étapes, 2 pers.)')
        if [(a, round(b, 3)) for a, b in nu['ingredients']] != [(a, b) for a, b in q1]:
            avertir(f'Recette {n} : ingrédients/quantités de Nutrition ≠ Recettes v3')
        nut_brut[rid(n)] = nu['acc']
        subs = e.get('substitutions') or 'aucune'
        recettes[rid(n)] = {
            'id': rid(n),
            'num': n,
            'nom': v['nom'],
            'source': e.get('source', ''),
            'source_url': (re.search(r'https?://\S+', e.get('source', '')) or [None])[0],
            'temps': e.get('temps', ''),
            'ustensiles': [u.strip() for u in re.split(r'\s\+\s', e.get('ustensiles', '')) if u.strip()],
            'substitutions': subs,
            'notes_substitution_v3': v['notes'],
            'nutrition_1portion': nu['total'],
            'ingredients_livres': v['ingredients'],
            'etapes': e['etapes'],
            'etapes_portions': 2,
        }

    # ---- semaines ----
    semaines = {}
    for k in sorted(sommaire, key=lambda s: int(s[1:])):
        som = sommaire[k]
        if k not in wb.sheetnames:
            avertir(f'Feuille {k} absente'); continue
        titre_s, planning, courses, weekend, placard = lire_semaine(wb[k], recettes, options_par_nom)
        planif = [p['recette_id'] for p in planning if p['recette_id']]
        if sorted(planif) != sorted(som['recettes']):
            avertir(f'{k} : recettes du planning {planif} ≠ Sommaire {som["recettes"]}')
        manquantes = [r for r in som['recettes'] if r not in recettes]
        if manquantes:
            avertir(f'{k} : recettes inconnues {manquantes}')

        tot_courses = round(sum(c['total'] for c in courses), 2)
        tot_placard = round(sum(p['total'] for p in placard), 2)
        tot_weekend = round(sum(w['total'] for w in weekend), 2)
        total = round(tot_courses + tot_placard + tot_weekend, 2)
        gachis = round(sum(c['gachis_eur'] for c in courses), 2)
        stock = round(sum(c['stock_eur'] for c in courses), 2)

        # contrôle vs valeurs littérales du Sommaire
        for nom, calc, lu in (('courses', tot_courses, som['courses']), ('placard', tot_placard, som['placard']),
                              ('week-end', tot_weekend, som['weekend']), ('total', total, som['total']),
                              ('gâchis', gachis, som['gachis']), ('stock', stock, som['stock'])):
            if lu is not None and abs(calc - lu) > 0.011:
                avertir(f'{k} : {nom} recalculé {calc} ≠ Sommaire {round(lu, 2)}')

        # moyenne sur les valeurs NON arrondies (comme le classeur)
        brutes = [nut_brut[r] for r in som['recettes'] if r in nut_brut]
        moy = {c: round(sum(b[i] for b in brutes) / len(brutes), 1) for i, c in enumerate(CHAMPS_NUT)} if brutes else {}
        moy['kcal'] = int(round(sum(b[0] for b in brutes) / len(brutes))) if brutes else 0

        def lien(nom):
            o = options_par_nom.get(nom)
            return {'produit': nom, 'option_id': o['id'] if o else None}

        semaines[k] = {
            'id': k,
            'titre': titre_s,
            'total_eur': total,
            'marge_eur': round(BUDGET - total, 2),
            'budget_max': BUDGET,
            'cout_par_portion': round(total / REPAS_PAR_SEMAINE, 2),
            'recettes_semaine': som['recettes'],
            'plat_weekend': lien(som['plat']),
            'pizza': lien(som['pizza']),
            'couts': {
                'courses_recettes': tot_courses,
                'placard': tot_placard,
                'weekend': tot_weekend,
                'total': total,
                'marge': round(BUDGET - total, 2),
                'gachis': gachis,
                'stock': stock,
            },
            'nutrition_moyenne': moy,
            'planning': planning,
            'courses': courses,
            'weekend': weekend,
            'placard': placard,
            'total_courses_recettes': tot_courses,
            'total_panier': round(tot_courses + tot_weekend, 2),
            'total_placard': tot_placard,
        }

    # ---- référentiel produits (export CSV « ingrédients ») ----
    ref = {}
    for k, w in semaines.items():
        for c in w['courses']:
            ref.setdefault(c['produit'], {
                'nom_chronodrive': c['produit'], 'rayon': c['rayon'], 'prix': c['prix_unitaire'],
                'contenu': arrondi(c['contenu_achete'] / c['qte_acheter'], 2) if c['contenu_achete'] and c['qte_acheter'] else c['contenu_achete'],
                'unite': c['unite'], 'type': 'course'})
        for p in w['placard']:
            ref.setdefault(p['produit'], {
                'nom_chronodrive': p['produit'], 'rayon': p['rayon'], 'prix': p['prix_format'],
                'contenu': p['contenu_format'], 'unite': p['unite'], 'type': 'placard'})
    for o in options:
        ref.setdefault(o['produit'], {
            'nom_chronodrive': o['produit'], 'rayon': 'Surgelé (week-end)', 'prix': o['prix'],
            'contenu': o['poids_g'], 'unite': 'g', 'type': 'weekend'})
    ref = dict(sorted(ref.items(), key=lambda kv: (kv[1]['type'], kv[1]['rayon'], kv[0])))

    # ---- contrôle nutrition vs valeurs en cache du classeur ----
    ws_cache = cache['Nutrition']
    for n, nu in sorted(nut.items()):
        li = nu.get('ligne_total')
        if not li:
            continue
        lu = ws_cache.cell(row=li, column=8).value
        if isinstance(lu, (int, float)) and abs(lu - nu['total']['kcal']) > 1:
            avertir(f'Recette {n} : kcal recalculé {nu["total"]["kcal"]} ≠ cache {round(lu, 1)}')
    som_cache = cache['Sommaire']
    for row in som_cache.iter_rows(values_only=True):
        k = txt(row[0])
        if k in semaines and isinstance(row[11], (int, float)):
            if abs(row[11] - semaines[k]['nutrition_moyenne']['kcal']) > 1:
                avertir(f'{k} : kcal moyen recalculé {semaines[k]["nutrition_moyenne"]["kcal"]} ≠ cache Sommaire {row[11]}')

    with open(chemin, 'rb') as f:
        empreinte = hashlib.sha256(f.read()).hexdigest()
    mtime = date.fromtimestamp(os.path.getmtime(chemin)).isoformat()

    return {
        'meta': {
            'version': 'V3.0',
            'genere_le': mtime,
            'source': os.path.basename(chemin),
            'source_sha256': empreinte,
            'titre': titre,
            'description': f'MealWeek V3 — {len(recettes)} recettes, {len(semaines)} semaines-types indépendantes '
                           f'(S1-S{len(semaines)}), Chronodrive Compiègne, retrait dimanche.',
            'modele': modele,
            'regles': {
                'portions': '1 portion = 1 personne. Recettes cuisinées x2 (dîner + midi lendemain). '
                            'Quantités des ingrédients stockées pour 1 portion ; les étapes sont rédigées pour 2 portions.',
                'semaine_type': f"L'utilisateur choisit 1 semaine-type parmi S1-S{len(semaines)} ; "
                                'la dernière sélectionnée est retenue et persistée.',
                'courses': "Les courses affichées = liste 'courses' de la semaine retenue (produits Chronodrive, "
                           'formats entiers, prix relevés). Retirer une recette retire ses ingrédients.',
            },
            'methode': methode,
        },
        'recipes': recettes,
        'ingredients_ref': ref,
        'options_weekend': options,
        'weeks': semaines,
    }


def main():
    chemin = sys.argv[1] if len(sys.argv) > 1 else XLSX_DEFAUT
    if not os.path.exists(chemin):
        sys.exit(f'Classeur introuvable : {chemin}')
    data = construire(chemin)
    texte = json.dumps(data, ensure_ascii=False, indent=2) + '\n'
    ancien = open(SORTIE, encoding='utf-8').read() if os.path.exists(SORTIE) else ''
    if ancien == texte:
        print('Aucun changement : mealweek_data.json est déjà synchronisé avec le classeur.')
    else:
        with open(SORTIE, 'w', encoding='utf-8') as f:
            f.write(texte)
        print(f'Écrit : {os.path.relpath(SORTIE, RACINE)}')
    print(f"{len(data['recipes'])} recettes · {len(data['weeks'])} semaines · "
          f"{len(data['options_weekend'])} options week-end · {len(data['ingredients_ref'])} produits")
    if avertissements:
        print(f'\n{len(avertissements)} contrôle(s) à regarder :')
        for a in avertissements:
            print('  - ' + a)
    else:
        print('Contrôles de cohérence : tous OK.')


if __name__ == '__main__':
    main()
